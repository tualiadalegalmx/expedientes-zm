"""
Genera el SQL de carga inicial a partir de la hoja "Plazos" y "Expedientes"
del archivo carga_inicial_*.xlsx.

Uso:
    python3 scripts/generar_importacion.py datos/carga_inicial_amparos_GrupoMAS.xlsx > datos/importacion.sql

Luego, en Supabase → SQL Editor, pegar y ejecutar datos/importacion.sql.
Requiere que el usuario responsable (por nombre, columna "Abogado responsable")
ya exista en la tabla perfiles.

IMPORTANTE: la carpeta datos/ contiene información de clientes y no se sube a GitHub.
"""
import sys
import datetime as dt
import openpyxl

ANIO = 2026


def q(v):
    if v is None or v == "":
        return "null"
    if isinstance(v, (dt.date, dt.datetime)):
        return f"'{v:%Y-%m-%d}'"
    if isinstance(v, (int, float)):
        return str(int(v))
    return "'" + str(v).replace("'", "''").strip() + "'"


def fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    return v


def main(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    out = ["-- Carga inicial generada automáticamente. Ejecutar una sola vez.", "begin;", ""]

    # Expedientes ---------------------------------------------------------
    we = wb["Expedientes"]
    hdr = [c.value for c in we[1]]
    ids = {}
    for fila in we.iter_rows(min_row=2, values_only=True):
        r = dict(zip(hdr, fila))
        if not r["Número de juicio"]:
            continue
        num = str(r["Número de juicio"]).strip()
        juz = (r["Juzgado o sala"] or "Por definir").strip()
        autoridad = r["Autoridad ante quien se promueve"] or (
            juz if juz == "Por definir" else f"{juz} en el Estado de Veracruz")
        resp = r["Abogado responsable"] or "Uriel"
        etapa = r["Etapa procesal"] or "Por definir"
        var = "e_" + str(len(ids))
        ids[num] = var
        out.append(f"""insert into expedientes (numero, juzgado, autoridad, lugar, cliente, contraparte, materia_id,
  tipo_juicio, responsable_id, etapa_id, estado, created_by, notas)
select {q(num)}, {q(juz)}, {q(autoridad)}, {q(r['Lugar'])}, {q(r['Cliente'])}, {q(r['Contraparte'])},
  (select id from materias where nombre = {q(r['Materia'])}),
  {q(r['Tipo de juicio o trámite'] or 'Juicio de amparo')},
  (select id from perfiles where nombre = {q(resp)}),
  (select id from etapas where nombre = {q(etapa)}),
  'activo', (select id from perfiles where nombre = {q(resp)}),
  'Importado del control de seguimiento de acuerdos (Excel).';""")
    out.append("")

    # Acuerdos, plazos y promociones -------------------------------------
    wp = wb["Plazos"]
    hdr = [c.value for c in wp[1]]
    n = 0
    for fila in wp.iter_rows(min_row=2, values_only=True):
        r = dict(zip(hdr, fila))
        num = r["Número de juicio"]
        if not num:
            continue
        num = str(num).strip()
        noti = fecha(r["Fecha de notificación"])
        venc = fecha(r["Vencimiento (Excel)"])
        pres = fecha(r["Fecha de presentación"])
        acuerdo = r["Actuación / acuerdo"] or "Acuerdo (sin descripción en el Excel)"
        notas = r["Notas"]
        cant = r["Plazo (cantidad)"]
        unidad = "horas" if r["Unidad"] == "Horas" else "dias"
        obs = []
        if pres and pres.year != ANIO:
            obs.append(f"año de presentación corregido de {pres.year} a {ANIO}")
            pres = pres.replace(year=ANIO)
        if venc and noti and venc < noti:
            obs.append(f"vencimiento del Excel ({venc:%d/%m/%Y}) era anterior a la notificación; revisar")
        texto_obs = (" [Importación: " + "; ".join(obs) + "]") if obs else ""
        exp = f"(select id from expedientes where numero = {q(num)})"
        autor = f"(select created_by from expedientes where numero = {q(num)})"
        n += 1
        a, p, pr = f"a{n}", f"p{n}", f"pr{n}"
        genera = "true" if cant else "false"
        desc = acuerdo + (f"\nNotas: {notas}" if notas else "") + texto_obs
        out.append(f"-- Fila {r['Fila original']} del Excel · {num}")
        out.append(f"""with {a} as (
  insert into bitacora (expediente_id, tipo, fecha, descripcion, genera_plazo, created_by)
  values ({exp}, 'acuerdo', {q(noti)}, {q(desc)}, {genera}, {autor}) returning id)""")
        if pres:
            out[-1] += f""",
{pr} as (
  insert into bitacora (expediente_id, tipo, fecha, descripcion, created_by)
  values ({exp}, 'promocion', {q(pres)}, {q('Promoción en atención a: ' + acuerdo[:200])}, {autor}) returning id)"""
        if cant and venc and noti:
            estado = "'cumplido'" if pres else "'abierto'"
            cumplido_con = f"(select id from {pr})" if pres else "null"
            out[-1] += f"""
insert into plazos (expediente_id, acuerdo_id, fecha_surte_efectos, cantidad, unidad, tipo_dias,
  actuacion, vencimiento, estado, cumplido_con, cumplido_fecha, folio, created_by)
select {exp}, (select id from {a}), {q(noti)}, {int(cant)}, '{unidad}', 'habiles',
  {q(acuerdo[:300])}, {q(venc)}, {estado}, {cumplido_con}, {q(pres)}, {q(r['Folio'])}, {autor};"""
        else:
            out[-1] += f"\nselect 1 from {a}" + (f", {pr}" if pres else "") + ";"
        out.append("")

    # Último movimiento real (no la fecha de importación)
    out.append("""update expedientes e set ultimo_movimiento = coalesce(
  (select max(b.fecha)::timestamptz from bitacora b where b.expediente_id = e.id), e.ultimo_movimiento);""")
    out.append("commit;")
    print("\n".join(out))
    print(f"-- {len(ids)} expedientes, {n} acuerdos", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1])
